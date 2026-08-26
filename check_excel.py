import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
wb = openpyxl.load_workbook('Cube_Bonus Potential Cube (Weapon).xlsx', read_only=True)
sheet_name = wb.sheetnames[0]
ws = wb[sheet_name]
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 100:
        print(f'{i}: {row}')